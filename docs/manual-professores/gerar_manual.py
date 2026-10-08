from pathlib import Path
import json,math
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
ROOT=Path(__file__).resolve().parent; META=json.loads((ROOT/'capturas.json').read_text()); PAGES=[]
def add(title,purpose,image,steps,tip,phase='PASSO A PASSO'):
 PAGES.append(dict(title=title,purpose=purpose,image=image,steps=steps,tip=tip,phase=phase))
def S(heading,explanation,**target):return dict(label=heading,explanation=explanation,target=target)
add('Manual do professor','APP Monitor • Guia ilustrado para aprender e consultar durante o uso.', '10-painel-professor',[
 S('Organize a aula','Aprenda a solicitar os aparelhos, acompanhar o pedido e registrar a utilização pela turma.'),
 S('Complete o empréstimo','Veja como devolver os dispositivos e assinar o relatório depois da conferência.'),
 S('Siga as setas','Os números nas imagens indicam os campos e botões explicados na mesma página.')], 'As telas usam pessoas e aparelhos fictícios. Endereço do app: app-monitor-interface.onrender.com','BEM-VINDO')
add('Como consultar este manual','Você pode ler na sequência ou ir direto à tarefa de que precisa.',None,[
 S('Primeiro acesso · páginas 3 a 7','Entrar, criar a conta de professor e trocar uma senha provisória.'),
 S('Preparar a aula · páginas 8 a 12','Conhecer o painel, solicitar os aparelhos e acompanhar a aprovação.'),
 S('Usar e devolver · páginas 13 a 21','Retirar, associar alunos, pedir extras, devolver, conferir e assinar.'),
 S('Consultar e resolver · páginas 22 a 28','Histórico, PDF, notificações, senha, celular e respostas às dúvidas mais comuns.')], 'Nos exemplos, TAB-001 identifica um aparelho fictício. Use os dados da sua própria aula.','ROTEIRO')
add('Entrar com sua conta','Abra https://app-monitor-interface.onrender.com no navegador.', '01-login',[
 S('Digite seu e-mail','Use o e-mail informado no cadastro. Confira se não há espaços extras ou erro de digitação.',name='email'),
 S('Digite sua senha','Use a senha da sua conta do App Monitor. Se a escola lhe entregou uma senha provisória, use-a neste primeiro acesso.',name='password'),
 S('Clique em Entrar','Depois de entrar, confira se o seu nome aparece no cabeçalho e se o perfil está como Professor.',text='Entrar')], 'Esqueceu a senha? Solicite a redefinição ao administrador da escola. Não crie outra conta para resolver isso.')
add('Ainda não tem uma conta?','O cadastro público está disponível somente para professores.', '01-login',[
 S('Abra o cadastro de professor','Na tela de entrada, clique em Sou professor — criar minha conta. O aplicativo abrirá um formulário.',id='registerTeacher'),
 S('Já tem uma conta criada pela escola?','Nesse caso, preencha o e-mail e a senha recebidos e use Entrar. Não é necessário se cadastrar novamente.',text='Entrar')], 'Se aparecer que o e-mail já está cadastrado, procure o administrador para recuperar o acesso.')
add('Cadastro: informe seus dados','Use informações que permitam à escola identificar você.', '02-cadastro-professor',[
 S('Nome completo','Digite seu nome inteiro, com nome e sobrenomes. Evite colocar somente o primeiro nome ou um apelido.',name='name'),
 S('E-mail','Informe o endereço que você utilizará para entrar no app. Leia novamente antes de continuar: um erro aqui pode dificultar o acesso.',name='email')], 'Exemplo fictício: Professora Ana Silva. Cada professor deve utilizar sua própria conta.')
add('Cadastro: crie sua senha','Conclua o formulário depois de preencher nome completo e e-mail.', '02-cadastro-professor',[
 S('Escolha uma senha','Digite uma senha com pelo menos 12 caracteres. Use uma senha que consiga guardar com segurança.',name='password'),
 S('Repita a mesma senha','Digite novamente, exatamente como no campo anterior. Os dois valores precisam coincidir.',name='confirmation'),
 S('Crie a conta','Clique em Criar conta de professor. Quando o cadastro for concluído, o aplicativo abrirá a área do professor.',text='Criar conta de professor')], 'O perfil criado aqui é Professor. O administrador da escola poderá consultar e gerenciar esse cadastro.')
add('Se recebeu uma senha provisória','Esta tela aparece quando a escola criou ou redefiniu sua conta.', '09-primeira-senha',[
 S('Informe a senha recebida','Em Senha provisória, digite a senha que o administrador lhe entregou.',name='current'),
 S('Defina sua nova senha','Crie uma senha de pelo menos 12 caracteres. No campo seguinte, confirme a mesma senha.',name='password'),
 S('Salve para continuar','Clique em Salvar minha senha. A troca precisa ser concluída antes de usar as funções do aplicativo.',text='Salvar minha senha')], 'Depois da troca, use a nova senha nos próximos acessos. A senha provisória deixa de ser a sua senha de entrada.')
add('Conheça a tela inicial','A Visão geral reúne seus pedidos e suas movimentações.', '10-painel-professor',[
 S('Agendamentos','Use este menu para acompanhar os pedidos de aparelhos feitos para suas aulas.',page='appointments'),
 S('Meus dispositivos','Abra este menu para ver os aparelhos que a equipe já liberou para você.',page='loans'),
 S('Notificações','O sino reúne avisos do aplicativo. Consulte-o para acompanhar mudanças e pendências.',id='notifications')], 'Sua conta de professor acompanha os seus próprios registros. Inventário e gestão de usuários são funções da equipe.')
add('Comece um novo pedido','Tenha em mãos a data, o horário, a turma e a quantidade necessária.', '10-painel-professor',[
 S('Clique em Novo agendamento','Na Visão geral, use o botão + Novo agendamento. Ele abre o formulário do pedido.',text='＋ Novo agendamento'),
 S('Também acompanhe pelo menu','O menu Agendamentos permite consultar os pedidos já registrados e iniciar um novo quando necessário.',page='appointments')], 'Criar o pedido ainda não significa que os aparelhos foram entregues: ele seguirá para aprovação da equipe.')
add('Preencha data, horário e turma','No formulário aberto, identifique a aula que usará os aparelhos.', '11-novo-agendamento',[
 S('Data da aula','Escolha a data correta. O aplicativo permite agendar entre hoje e os próximos 14 dias.',name='date'),
 S('Horário','Informe o horário previsto para a utilização dos equipamentos.',name='time'),
 S('Turma','Digite a identificação da turma. Exemplo: 8º Ano B. Use o padrão de identificação adotado pela escola.',name='class_name')], 'Confira esses três campos antes de salvar, principalmente quando fizer pedidos para mais de uma turma.')
add('Escolha os aparelhos e salve','O pedido usa o tipo do equipamento, não o modelo ou fabricante.', '11-novo-agendamento',[
 S('Tipo de equipamento','Escolha Tablet, Notebook, Chromebook ou Celular, conforme a necessidade da aula e a orientação da escola.',name='type'),
 S('Quantidade','Informe a quantidade necessária: de 1 a 25 aparelhos por pedido. Confira o número antes de enviar.',name='quantity'),
 S('Salvar agendamento','Clique neste botão. Após salvar, abra Agendamentos e confira se o pedido aparece na lista.',text='Salvar agendamento')], 'Exemplo: para uma atividade em duplas, calcule a quantidade necessária para a turma antes de enviar o pedido.')
add('Acompanhe e ajuste o pedido','Abra Agendamentos para consultar a situação da sua solicitação.', '12-pedido-pendente',[
 S('Entenda o status','Pendente significa que a equipe ainda precisa analisar o pedido. Aprovado significa que ele foi aprovado; combine a retirada.',text='Pendente'),
 S('Editar','Antes da liberação, use Editar se precisar corrigir os dados. Ao salvar uma alteração, o pedido volta para aprovação.',text='Editar'),
 S('Cancelar','Se não for mais usar os aparelhos, cancele o pedido enquanto essa ação estiver disponível. Confira se escolheu a solicitação certa.',text='Cancelar')], 'A liberação dos dispositivos ocorre no dia agendado e é registrada pela equipe responsável.')
add('Depois de retirar os aparelhos','A equipe precisa registrar a liberação para que o empréstimo apareça.', '26-meus-dispositivos',[
 S('Abra Meus dispositivos','Acesse esse menu depois da entrega dos aparelhos.',page='loans'),
 S('Localize a movimentação','Confira turma, quantidade e situação. Em uso indica que o empréstimo está em andamento.',text='Em uso'),
 S('Abra os detalhes','Clique em Abrir movimentação para ver a lista de aparelhos e registrar quem está utilizando cada um.',text='Abrir movimentação')], 'Se a movimentação não aparecer, confirme com a equipe se a liberação foi concluída na sua conta e turma.')
add('Associe os aparelhos aos alunos','Cada linha mostra o identificador de um aparelho entregue.', '16-uso-e-devolucao',[
 S('Preencha os alunos','Ao lado de cada dispositivo, informe o aluno que o utilizou, conforme a orientação da escola. Repita nas demais linhas.',label='Aluno do dispositivo TAB-001'),
 S('Salve as associações','Clique em Salvar associações para registrar os nomes durante a aula. Depois, você poderá continuar usando a movimentação.',id='saveStudents')], 'Confira a etiqueta física para não trocar os alunos entre os aparelhos. Todas as associações devem estar preenchidas ao devolver.')
add('Precisa de aparelhos extras?','A solicitação de extras acontece dentro de um empréstimo em uso.', '16-uso-e-devolucao',[
 S('Abra a solicitação','Na movimentação, clique em Solicitar extras. Será aberto um formulário de quantidade.',id='requestExtras'),
 S('Mantenha o registro dos alunos','Se preencher ou corrigir os alunos durante a aula, use Salvar associações para guardar as alterações.',id='saveStudents')], 'É permitido um único pedido de até 5 aparelhos extras por aula. A equipe ainda precisa efetuar a liberação dos extras.')
add('Envie o pedido de extras','Escolha a quantidade adicional com atenção.', '21-extras',[
 S('Quantidade extra','Informe de 1 a 5 aparelhos adicionais. Essa quantidade é acrescentada ao pedido original, quando liberada pela equipe.',name='quantity'),
 S('Enviar pedido','Clique em Enviar pedido e avise a equipe responsável pela entrega. Confira os novos aparelhos na movimentação após a liberação.',text='Enviar pedido')], 'Não registre aparelhos recebidos por fora do fluxo. Peça à equipe para registrar a entrega no App Monitor.')
add('Registre e envie a devolução','Ao concluir a aula, confira todos os aparelhos e devolva-os à equipe.', '16-uso-e-devolucao',[
 S('Escreva o relato','Descreva como terminou a utilização e informe qualquer ocorrência. Exemplo: aula concluída e aparelhos devolvidos sem ocorrências.',id='returnReport'),
 S('Envie para conferência','Confira os alunos e clique em Enviar devolução. O aplicativo salva as associações e encaminha a movimentação para conferência.',id='submitReturn')], 'Se houve dano, falha ou falta de um aparelho, registre o ocorrido e avise a equipe. O relato deve ter pelo menos 5 caracteres.')
add('Aguarde a conferência da equipe','Enviar a devolução é uma etapa; o retorno físico ainda será conferido.', '27-devolucao-pendente',[
 S('Consulte Meus dispositivos','Acompanhe a movimentação enquanto a equipe verifica os aparelhos recebidos.',page='loans'),
 S('Observe o status de retorno','Conferir retorno indica que os aparelhos aguardam a conferência registrada pela equipe.',text='Conferir retorno')], 'Depois da conferência, será necessário revisar e assinar o relatório. Fique atento às notificações.')
add('Abra o relatório para assinar','Acesse o relatório depois que a equipe concluir a conferência.', '28-lista-relatorios',[
 S('Relatórios e assinaturas','Abra este menu para encontrar os relatórios e as assinaturas pendentes.',page='reports'),
 S('Ver relatório','Localize a movimentação da sua turma e clique em Ver relatório.',text='Ver relatório')], 'Uma assinatura sua pendente pode impedir novos pedidos. Resolva a pendência após conferir as informações do relatório.')
add('Leia antes de assinar','Confira se o relatório corresponde ao que aconteceu na aula.', '18-assinatura',[
 S('Identificação da movimentação','Confira professor, turma, situação e datas. Leia também a lista dos aparelhos e o relato registrado.',cls='report-meta'),
 S('Assinaturas','Observe quais assinaturas já foram registradas e quais estão pendentes. Você assina pela sua própria conta de professor.',cls='signature-grid')], 'Se encontrar uma divergência, procure a equipe responsável antes de confirmar sua assinatura.')
add('Desenhe e confirme a assinatura','Role a janela do relatório até a área da sua assinatura.', '22-assinar-campo',[
 S('Desenhe neste espaço','No celular, use o dedo. No computador, clique e arraste com o mouse para desenhar sua assinatura.',id='signature'),
 S('Limpar, se precisar refazer','Se o desenho ficar incorreto, clique em Limpar e faça a assinatura novamente.',id='clearSignature'),
 S('Confirme o registro','Quando terminar, clique em Confirmar minha assinatura. Aguarde a confirmação antes de sair.',id='signReport')], 'As assinaturas são registros manuscritos feitos por contas autenticadas; não são assinaturas certificadas ICP-Brasil.')
add('Consulte seus registros concluídos','Depois das assinaturas necessárias, a movimentação fica disponível no histórico.', '25-historico',[
 S('Abra Histórico','Esse menu reúne as movimentações concluídas que você pode consultar.',page='history'),
 S('Pesquise a movimentação','Use o campo de pesquisa para filtrar pela turma ou pelo professor, conforme os registros disponíveis para sua conta.',id='historySearch'),
 S('Visualizar / PDF','Clique no botão da linha desejada para abrir novamente o relatório.',text='Visualizar / PDF')], 'Se o empréstimo ainda não estiver no histórico, verifique se falta conferência ou alguma assinatura.')
add('Guarde uma cópia do relatório','Abra o relatório pelo Histórico antes de imprimir ou salvar.', '19-relatorio',[
 S('Imprimir / salvar PDF','Clique neste botão para abrir a janela de impressão do navegador.',id='printReport'),
 S('Confira as assinaturas','Antes de guardar a cópia, verifique se o relatório é o desejado e se contém os registros esperados.',cls='signature-grid')], 'Na janela do navegador, escolha Salvar como PDF (ou opção equivalente), indique a pasta e confirme Salvar.')
add('Consulte as notificações','Clique no sino, no cabeçalho do aplicativo, para abrir os avisos.', '23-notificacoes',[
 S('Leia e abra o aviso','Confira a mensagem e clique em Abrir para ir à área relacionada à notificação.',text='Abrir'),
 S('Concluir a notificação','Use Concluir quando tiver tratado aquele aviso. Esse botão conclui o aviso, não substitui as ações do empréstimo.',text='Concluir')], 'Marcar um aviso como concluído não assina relatórios, não devolve aparelhos e não cancela agendamentos.')
add('Troque sua senha quando precisar','Depois de entrar, clique em Conta no cabeçalho.', '24-conta',[
 S('Senha atual','Digite a senha que você está usando para entrar no App Monitor.',name='current'),
 S('Nova senha','Escolha outra senha com pelo menos 12 caracteres. Confira a digitação antes de salvar.',name='password'),
 S('Salvar senha','Clique em Salvar senha. Passe a usar a nova senha nos próximos acessos.',text='Salvar senha')], 'Se não souber a senha atual e não conseguir entrar, peça ao administrador uma nova senha provisória.')
add('Use também pelo celular','Abra o mesmo endereço no navegador e entre com sua própria conta.', '20-celular',[
 S('Acompanhe os pedidos','Use Agendamentos para consultar as solicitações mesmo estando fora do computador.',page='appointments'),
 S('Acesse os empréstimos','Em Meus dispositivos, você pode preencher os alunos e registrar a devolução.',page='loans'),
 S('Saia ao terminar','Em um aparelho compartilhado, clique em Sair para encerrar sua sessão.',id='logout')], 'Se desejar, use a opção do navegador para adicionar um atalho à tela inicial. As funções continuam disponíveis pelo site.')
add('Dúvidas e soluções rápidas','Antes de fazer outro cadastro ou repetir uma ação, confira estas orientações.',None,[
 S('Não consigo entrar','Revise o e-mail e a senha. Se esqueceu a senha, peça a redefinição ao administrador. Se a conta estiver desativada, procure a escola.'),
 S('O site demora a abrir','A primeira conexão pode demorar enquanto o servidor gratuito é ativado. Aguarde a conexão e siga a orientação de tentar novamente, se aparecer.'),
 S('Não consigo devolver ou pedir novamente','Preencha todos os alunos e o relato para devolver. Para novos pedidos, verifique se existe relatório pendente da sua assinatura.'),
 S('Há erro no pedido ou problema com um aparelho','Antes da liberação, use Editar quando disponível. Durante o uso, comunique a equipe e registre a ocorrência no relato.')], 'Ao pedir ajuda, informe a turma, a data e a mensagem exibida. Nunca envie sua senha.','APOIO')
add('Checklist do professor','Use esta página como lembrete antes, durante e depois da aula.',None,[
 S('Antes da aula','Entrar na própria conta → registrar o pedido → acompanhar a aprovação → combinar a retirada.'),
 S('Durante a aula','Conferir os aparelhos recebidos → preencher e salvar os alunos → comunicar problemas e solicitar extras, se necessário.'),
 S('Ao terminar','Conferir e entregar os aparelhos → escrever o relato → enviar a devolução → aguardar a conferência.'),
 S('Para concluir','Ler o relatório → assinar → consultar o Histórico e salvar o PDF quando precisar.')], 'Acesse: https://app-monitor-interface.onrender.com','CONSULTA RÁPIDA')
ROOT.joinpath('roteiro.json').write_text(json.dumps(PAGES,ensure_ascii=False,indent=2),encoding='utf8')
W,H=595.276,841.89; NAVY='#142843'; BLUE='#206DE3'; GREEN='#109C79'; MUTED='#506078'; BG='#F4F7FC'; WHITE='#FFFFFF'
pdfmetrics.registerFont(TTFont('Manual','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'));pdfmetrics.registerFont(TTFont('ManualBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
c=canvas.Canvas(str(ROOT/'APP-Monitor-Manual-do-Professor.pdf'),pagesize=(W,H));c.setTitle('APP Monitor — Manual ilustrado do professor');c.setAuthor('APP Monitor')
prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H);prs.core_properties.title='APP Monitor — Manual ilustrado do professor';prs.core_properties.author='APP Monitor'
current=None

def rgb(v):return RGBColor.from_string(v.strip('#'))
def box(x,y,w,h,fill,line=None,radius=False):
 global current
 shape=current.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill)
 if line:shape.line.color.rgb=rgb(line);shape.line.width=Pt(1)
 else:shape.line.fill.background()
 c.setFillColor(fill);c.setStrokeColor(line or fill)
 if radius:c.roundRect(x,H-y-h,w,h,8,fill=1,stroke=bool(line))
 else:c.rect(x,H-y-h,w,h,fill=1,stroke=bool(line))
def lines(text,size,width,bold=False):
 font='ManualBold' if bold else 'Manual';out=[];line=''
 for word in text.split():
  trial=(line+' '+word).strip()
  if line and pdfmetrics.stringWidth(trial,font,size)>width:out.append(line);line=word
  else:line=trial
 if line:out.append(line)
 return out

def text(x,y,w,content,size=12,fill=NAVY,bold=False):
 ls=lines(content,size,w,bold);height=len(ls)*size*1.3
 t=current.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(height+4));tf=t.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
 for i,l in enumerate(ls):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=l;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(fill);p.space_after=Pt(0);p.line_spacing=Pt(size*1.3)
 c.setFont('ManualBold' if bold else 'Manual',size);c.setFillColor(fill)
 for i,l in enumerate(ls):c.drawString(x,H-y-size-i*size*1.3,l)
 return height

def circle(x,y,n,r=9):
 sh=current.shapes.add_shape(MSO_SHAPE.OVAL,Pt(x-r),Pt(y-r),Pt(r*2),Pt(r*2));sh.fill.solid();sh.fill.fore_color.rgb=rgb(BLUE);sh.line.fill.background()
 c.setFillColor(BLUE);c.circle(x,H-y,r,fill=1,stroke=0)
 text(x-r,y-r+1,r*2,str(n),11,WHITE,True)
 # Center PDF digit, replacing previous left-biased glyph with a clean circle.
 c.setFillColor(BLUE);c.circle(x,H-y,r,fill=1,stroke=0);c.setFillColor(WHITE);c.setFont('ManualBold',11);c.drawCentredString(x,H-y-3.8,str(n))
 tf=current.shapes[-1].text_frame
 from pptx.enum.text import PP_ALIGN
 tf.paragraphs[0].alignment=PP_ALIGN.CENTER

def arrow(x1,y1,x2,y2):
 sh=current.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2));sh.line.color.rgb=rgb(BLUE);sh.line.width=Pt(1.6)
 end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','med');end.set('len','med');sh.line._get_or_add_ln().append(end)
 c.setStrokeColor(BLUE);c.setFillColor(BLUE);c.setLineWidth(1.6);c.line(x1,H-y1,x2,H-y2)
 angle=math.atan2(y2-y1,x2-x1);p=c.beginPath();p.moveTo(x2,H-y2)
 for a in [angle+2.7,angle-2.7]:p.lineTo(x2+7*math.cos(a),H-(y2+7*math.sin(a)))
 p.close();c.drawPath(p,fill=1,stroke=0)

def outline(x,y,w,h):
 sh=current.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));sh.fill.background();sh.line.color.rgb=rgb(BLUE);sh.line.width=Pt(1.4)
 c.setStrokeColor(BLUE);c.setLineWidth(1.4);c.rect(x,H-y-h,w,h,fill=0,stroke=1)

def annotated_image(image,steps):
 path=ROOT/'imagens'/f'{image}.png';iw,ih=Image.open(path).size
 scale=min(455/iw,340/ih);pw,ph=iw*scale,ih*scale;px=(W-pw)/2;py=159+(340-ph)/2
 current.shapes.add_picture(str(path),Pt(px),Pt(py),width=Pt(pw),height=Pt(ph));c.drawImage(str(path),px,H-py-ph,width=pw,height=ph,mask="auto")
 targetmeta=META[image];sx=pw/targetmeta['width'];sy=ph/targetmeta['height'];used={'L':[],'R':[]}
 for n,step in enumerate(steps,1):
  criteria=step['target']
  if not criteria:continue
  found=[t for t in targetmeta['targets'] if all(t.get(k)==v for k,v in criteria.items())]
  if not found:raise RuntimeError(f'Target ausente: {image} {criteria}')
  target=found[0];x=px+target['x']*sx;y=py+target['y']*sy;w=target['w']*sx;h=target['h']*sy
  outline(x-1,y-1,w+2,h+2)
  side='L' if x+w/2<=px+pw/2 else 'R';cy=y+h/2
  for other in used[side]:
   if abs(cy-other)<25:cy=other+25
  cy=min(500,max(158,cy));used[side].append(cy)
  cx=px-23 if side=='L' else px+pw+23
  circle(cx,cy,n)
  arrow(cx+(10 if side=='L' else -10),cy,x-3 if side=='L' else x+w+3,y+h/2)
 text(45,512,W-90,'As setas numeradas correspondem às explicações abaixo.',9,MUTED)

for num,page in enumerate(PAGES,1):
 current=prs.slides.add_slide(prs.slide_layouts[6]);box(0,0,W,H,WHITE);box(0,0,W,7,BLUE)
 box(35,24,31,27,GREEN);text(39,28,26,'AM',12,WHITE,True);text(77,29,420,'APP MONITOR  /  MANUAL DO PROFESSOR',9,MUTED,True)
 c.bookmarkPage('page'+str(num));c.addOutlineEntry(page['title'],'page'+str(num),0,False)
 titleheight=text(35,67,W-70,page['title'],24,NAVY,True)
 purposeheight=text(35,74+titleheight,W-70,page['purpose'],11.3,MUTED)
 if page['image']:annotated_image(page['image'],page['steps']);yy=540
 else:yy=192
 for n,step in enumerate(page['steps'],1):
  circle(45,yy+8,n)
  head=text(64,yy,W-103,step['label'],13.1,NAVY,True)
  body=text(64,yy+head+4,W-103,step['explanation'],11.8,NAVY)
  yy+=head+body+18
 if yy>742:raise RuntimeError(f'Texto excede área útil: página {num} {yy}')
 box(35,749,W-70,55,BG,radius=True);text(47,758,W-94,page['tip'],10.1,MUTED)
 text(35,819,W-95,'Telas e dados de demonstração • Outubro de 2026',8,MUTED);text(W-66,817,40,f'{num:02d}/{len(PAGES)}',9,MUTED,True)
 if 'app-monitor-interface.onrender.com' in page['purpose']:
  c.linkURL('https://app-monitor-interface.onrender.com',(35,H-74-titleheight-purposeheight,W-35,H-74-titleheight),relative=0)
 if 'app-monitor-interface.onrender.com' in page['tip']:
  c.linkURL('https://app-monitor-interface.onrender.com',(35,35,W-35,95),relative=0)
 for shape in current.shapes:
  if shape.has_text_frame and 'app-monitor-interface.onrender.com' in shape.text:
   for paragraph in shape.text_frame.paragraphs:
    for run in paragraph.runs:run.hyperlink.address='https://app-monitor-interface.onrender.com'
 current.notes_slide.notes_text_frame.text=page['purpose']+'\n\n'+'\n\n'.join(f'{i}. {s["label"]}: {s["explanation"]}' for i,s in enumerate(page['steps'],1))+'\n\n'+page['tip']
 c.showPage()
c.save();prs.save(ROOT/'APP-Monitor-Manual-do-Professor.pptx')
print('Manual gerado:',len(PAGES),'páginas. PDF e PowerPoint editável com setas vetoriais.')
